"""Read every published register row; retain identical anonymous records."""

import hashlib
import json
from collections import Counter
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook

COLUMNS = [
    "location",
    "power",
    "self_protection",
    "ddt_average",
    "ddt_maximum",
    "burning_rate",
    "criticality",
    "connectors",
    "client_type",
    "users",
    "eens",
    "installation_type",
    "air_network",
    "circuit_queue",
    "network_km",
    "burned",
]
CATEGORICAL = ["client_type", "installation_type"]
BINARY = [
    "location",
    "self_protection",
    "criticality",
    "connectors",
    "air_network",
    "circuit_queue",
    "burned",
]


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_register(path: Path, year: int, expected_fields: list[str]) -> Any:
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        if book.sheetnames != ["Datos"]:
            raise ValueError("Unexpected sheets")
        iterator = book["Datos"].iter_rows(values_only=True)
        header = list(next(iterator))
        if header != expected_fields:
            raise ValueError("Unexpected source schema")
        rows = list(iterator)
    finally:
        book.close()
    if any(len(row) != len(COLUMNS) for row in rows):
        raise ValueError("Unexpected row width")
    frame = pd.DataFrame(rows, columns=COLUMNS)
    for col in COLUMNS:
        if col not in CATEGORICAL:
            frame[col] = pd.to_numeric(frame[col], errors="coerce")
    if not frame.burned.isin([0, 1]).all():
        raise ValueError("Missing or nonbinary annual outcome; do not invent labels")
    frame["record"] = [f"{year}:Datos:{i + 2}" for i in range(len(frame))]
    frame["year"] = year
    return frame


def canonical(value: Any) -> str | float | None:
    if pd.isna(value):
        return None
    if isinstance(value, (int, float, np.integer, np.floating)):
        return float(value)
    return str(value)


def profiles(frame: Any, columns: list[str]) -> list[str]:
    result = []
    for row in frame[columns].itertuples(index=False, name=None):
        encoded = json.dumps(
            [canonical(v) for v in row], ensure_ascii=True, separators=(",", ":"), allow_nan=False
        )
        result.append(hashlib.sha256(encoded.encode()).hexdigest())
    return result


def issue_masks(frame: Any) -> dict[str, Any]:
    masks = {}
    for col in COLUMNS:
        masks[f"missing_{col}"] = frame[col].isna()
    for col in BINARY:
        masks[f"invalid_binary_{col}"] = ~frame[col].isin([0, 1])
    for col in ("ddt_average", "ddt_maximum", "burning_rate", "users", "eens", "network_km"):
        masks[f"invalid_nonnegative_{col}"] = (frame[col] < 0) | ~np.isfinite(frame[col])
    masks["invalid_power"] = (frame.power <= 0) | ~np.isfinite(frame.power)
    masks["noninteger_users"] = frame.users.notna() & (frame.users % 1 != 0)
    masks["maximum_below_average_ddt"] = frame.ddt_maximum < frame.ddt_average
    masks["zero_users"] = frame.users == 0
    # Audit flag only: threshold is a research plausibility check, not a unit correction.
    masks["suspect_network_length_over_100_km"] = frame.network_km > 100
    return masks


def sanitise(frame: Any) -> Any:
    result = frame.copy()
    for col in COLUMNS:
        if col not in CATEGORICAL:
            result[col] = result[col].astype(float)
    for col in BINARY:
        result.loc[~result[col].isin([0, 1]), col] = np.nan
    for col in ("ddt_average", "ddt_maximum", "burning_rate", "users", "eens", "network_km"):
        result.loc[(result[col] < 0) | ~np.isfinite(result[col]), col] = np.nan
    result.loc[(result.power <= 0) | ~np.isfinite(result.power), "power"] = np.nan
    result.loc[result.users % 1 != 0, "users"] = np.nan
    result.loc[result.ddt_maximum < result.ddt_average, "ddt_maximum"] = np.nan
    return result


def audit(frame: Any, predictors: list[str]) -> dict[str, Any]:
    masks = issue_masks(frame)
    group = profiles(frame, predictors)
    frequencies = Counter(group)
    mixed = frame.assign(profile=group).groupby("profile").burned.nunique()
    return {
        "records": len(frame),
        "failures": int(frame.burned.sum()),
        "issue_counts": {key: int(value.sum()) for key, value in masks.items()},
        "rows_with_any_issue": int(np.logical_or.reduce(list(masks.values())).sum()),
        "exact_duplicate_excess": int(frame[COLUMNS].duplicated().sum()),
        "rows_in_exact_duplicate_sets": int(frame[COLUMNS].duplicated(keep=False).sum()),
        "distinct_predictor_profiles": len(frequencies),
        "repeated_profile_excess": len(frame) - len(frequencies),
        "rows_in_repeated_profiles": sum(n for n in frequencies.values() if n > 1),
        "profiles_with_mixed_outcomes": int((mixed > 1).sum()),
        "age_unavailable": len(frame),
        "health_unavailable": len(frame),
        "failure_dates_unavailable": len(frame),
        "asset_identity_unavailable": len(frame),
        "duplicate_policy": "Retain all: identical attributes do not prove duplicate anonymous assets",
        "numeric_ranges": {
            col: {"minimum": float(frame[col].min()), "maximum": float(frame[col].max())}
            for col in COLUMNS
            if col not in CATEGORICAL
        },
    }
